"""Extract text and table contents from common study-document formats."""
import csv
import io
import json
import re
import shutil
import zipfile
from html.parser import HTMLParser
from pathlib import Path

import fitz
from PIL import Image, ImageOps
import pytesseract

from nlp_engine import engine, TOPIC_FILES
from pdf_service import build_teaching_plan as _build_teaching_plan, chunk_text

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".pptx", ".txt", ".md", ".csv", ".xlsx",
    ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp",
    ".html", ".htm", ".json",
}
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 5000


def _clean(text):
    return re.sub(r"\n{3,}", "\n\n", str(text or "")).strip()


def _validate_office_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise ValueError("Office document contains too many embedded files")
        if sum(entry.file_size for entry in entries) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
            raise ValueError("Office document expands beyond the allowed processing size")


def _ocr_image(image_bytes):
    if not shutil.which("tesseract"):
        raise RuntimeError(
            "Image OCR requires the Tesseract OCR application. Install Tesseract and add it to PATH, then restart the backend."
        )
    with Image.open(io.BytesIO(image_bytes)) as image:
        image = image.convert("RGB")
        image.thumbnail((3000, 3000))
        image = ImageOps.autocontrast(image.convert("L"))
        return _clean(pytesseract.image_to_string(image))


def _image_text_or_empty(image_bytes):
    """OCR an embedded image if Tesseract is installed; preserve Office text otherwise."""
    if not shutil.which("tesseract"):
        return ""
    try:
        return _ocr_image(image_bytes)
    except Exception:
        return ""


def _extract_pdf(data):
    units = []
    with fitz.open(stream=data, filetype="pdf") as document:
        for index, page in enumerate(document, 1):
            text = _clean(page.get_text("text"))
            if not text and shutil.which("tesseract"):
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                text = _image_text_or_empty(pixmap.tobytes("png"))
            if shutil.which("tesseract"):
                image_text = []
                seen_xrefs = set()
                for image_info in page.get_images(full=True):
                    xref = image_info[0]
                    if xref in seen_xrefs:
                        continue
                    seen_xrefs.add(xref)
                    extracted = document.extract_image(xref)
                    recognized = _image_text_or_empty(extracted["image"])
                    if recognized:
                        image_text.append(recognized)
                if image_text:
                    text = _clean("\n\n".join([text, *image_text]))
            units.append({"page": index, "source_type": "page", "source_label": f"Page {index}", "text": text})
    return units


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.skip_depth += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.skip_depth:
            self.skip_depth -= 1
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip_depth:
            self.parts.append(data)


def _extract_docx(data):
    from docx import Document

    _validate_office_archive(data)
    document = Document(io.BytesIO(data))
    parts = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())
    for table_index, table in enumerate(document.tables, 1):
        rows = [" | ".join(_clean(cell.text) for cell in row.cells) for row in table.rows]
        if rows:
            parts.append(f"Table {table_index}:\n" + "\n".join(rows))
    units = [{"page": 1, "source_type": "document", "source_label": "Document text and tables", "text": _clean("\n\n".join(parts))}]
    if shutil.which("tesseract"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            image_index = 0
            for name in archive.namelist():
                if name.startswith("word/media/") and not name.endswith("/"):
                    image_index += 1
                    ocr_text = _image_text_or_empty(archive.read(name))
                    if ocr_text:
                        units.append({"page": len(units) + 1, "source_type": "image", "source_label": f"Embedded image {image_index}", "text": ocr_text})
    return units


def _extract_pptx(data):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    _validate_office_archive(data)
    presentation = Presentation(io.BytesIO(data))
    units = []
    for slide_index, slide in enumerate(presentation.slides, 1):
        parts = []
        image_index = 0
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = _clean(shape.text)
                if text:
                    parts.append(text)
            if shape.has_table:
                rows = [" | ".join(_clean(cell.text) for cell in row.cells) for row in shape.table.rows]
                parts.append("Table:\n" + "\n".join(rows))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_index += 1
                ocr_text = _image_text_or_empty(shape.image.blob)
                if ocr_text:
                    parts.append(f"Image {image_index} OCR:\n{ocr_text}")
        units.append({"page": slide_index, "source_type": "slide", "source_label": f"Slide {slide_index}", "text": _clean("\n\n".join(parts))})
    return units


def _extract_xlsx(data):
    from openpyxl import load_workbook

    _validate_office_archive(data)
    workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    units = []
    for sheet_index, sheet in enumerate(workbook.worksheets, 1):
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [f"Column {index}: {value}" for index, value in enumerate(row, 1) if value is not None and str(value).strip()]
            if cells:
                rows.append(" | ".join(cells))
        units.append({"page": sheet_index, "source_type": "sheet", "source_label": f"Sheet: {sheet.title}", "text": "\n".join(rows)})
        for image_index, image in enumerate(getattr(sheet, "_images", []), 1):
            ocr_text = _image_text_or_empty(image._data())
            if ocr_text:
                units.append({"page": len(units) + 1, "source_type": "image", "source_label": f"{sheet.title} image {image_index}", "text": ocr_text})
    workbook.close()
    return units


def extract_units(data, filename):
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Supported extensions: {allowed}")
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
        text = _ocr_image(data)
        return [{"page": 1, "source_type": "image", "source_label": "Image OCR", "text": text}]
    if extension == ".docx":
        return _extract_docx(data)
    if extension == ".pptx":
        return _extract_pptx(data)
    if extension == ".xlsx":
        return _extract_xlsx(data)
    decoded = data.decode("utf-8-sig", errors="replace")
    if extension == ".csv":
        rows = [" | ".join(cell.strip() for cell in row if cell.strip()) for row in csv.reader(io.StringIO(decoded))]
        text = "\n".join(row for row in rows if row)
        return [{"page": 1, "source_type": "table", "source_label": "CSV table", "text": text}]
    if extension in {".html", ".htm"}:
        parser = _HTMLText()
        parser.feed(decoded)
        decoded = "".join(parser.parts)
    elif extension == ".json":
        decoded = json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
    if extension == ".md":
        sections = re.split(r"(?m)(?=^#{1,6}\s+)", decoded)
        units = []
        for index, section in enumerate(sections, 1):
            text = _clean(section)
            if text:
                heading = next((line.lstrip("# ").strip() for line in text.splitlines() if line.startswith("#")), "")
                units.append({"page": index, "source_type": "section", "source_label": f"Section: {heading}" if heading else f"Section {index}", "text": text})
        return units or [{"page": 1, "source_type": "document", "source_label": "Text", "text": _clean(decoded)}]
    return [{"page": 1, "source_type": "document", "source_label": "Text", "text": _clean(decoded)}]


def process_document(data, filename):
    units = extract_units(data, filename)
    if not any(unit["text"] for unit in units):
        suffix = Path(filename).suffix.lower()
        if suffix in {".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"} and not shutil.which("tesseract"):
            raise ValueError("No selectable text was found. Scanned PDFs and images require Tesseract OCR to be installed and added to PATH.")
        raise ValueError("No readable text or table data was found in this file.")

    chunks = []
    for unit in units:
        for text in chunk_text(unit["text"]):
            if text.strip():
                chunks.append({"page": unit["page"], "source_type": unit["source_type"],
                               "source_label": unit["source_label"], "text": text})
    full_text = "\n\n".join(unit["text"] for unit in units)
    topic_key, topic_name = _detect_topic(full_text)
    if chunks:
        embeddings = engine.encode([chunk["text"] for chunk in chunks]).cpu().numpy().tolist()
        for chunk, embedding in zip(chunks, embeddings):
            chunk["embedding"] = embedding
    warnings = []
    if not shutil.which("tesseract"):
        warnings.append("Image text and scanned PDF pages require Tesseract OCR; selectable text and tables were extracted normally.")
    return {"pages": units, "num_pages": len(units), "chunks": chunks,
            "topic_key": topic_key, "topic_name": topic_name, "full_text": full_text[:4000],
            "warnings": warnings}


def _detect_topic(text):
    text = text.lower()
    scores = {topic: sum(1 for keyword in engine.keywords.get(topic, []) if keyword in text) for topic in TOPIC_FILES}
    best = max(scores, key=scores.get) if scores else None
    if best and scores[best] > 0:
        return best, engine.display_names.get(best)
    return None, "General"


def build_teaching_plan(chunks, topic_name):
    return _build_teaching_plan(chunks, topic_name)
