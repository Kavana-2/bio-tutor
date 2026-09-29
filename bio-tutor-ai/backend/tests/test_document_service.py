import io

import pytest

import document_service


def test_extracts_plain_text_and_markdown_sections():
    units = document_service.extract_units(
        b"Photosynthesis converts light energy into chemical energy.", "notes.txt"
    )
    assert units[0]["source_label"] == "Text"
    assert "chemical energy" in units[0]["text"]

    markdown = document_service.extract_units(
        b"# Photosynthesis\nLight reactions\n## Calvin cycle\nCarbon fixation", "notes.md"
    )
    assert len(markdown) == 2
    assert markdown[0]["source_label"] == "Section: Photosynthesis"


def test_extracts_csv_rows_as_table_text():
    content = "Topic,Definition\nPhotosynthesis,Plants convert light into chemical energy\n".encode()
    unit = document_service.extract_units(content, "biology.csv")[0]
    assert unit["source_type"] == "table"
    assert "Topic | Definition" in unit["text"]
    assert "Plants convert light into chemical energy" in unit["text"]


def test_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="Unsupported file type"):
        document_service.extract_units(b"data", "legacy.doc")


def test_extracts_docx_paragraphs_and_tables():
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Photosynthesis uses sunlight.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Input"
    table.cell(0, 1).text = "Carbon dioxide"
    stream = io.BytesIO()
    document.save(stream)

    units = document_service.extract_units(stream.getvalue(), "lesson.docx")
    text = "\n".join(unit["text"] for unit in units)
    assert "Photosynthesis uses sunlight" in text
    assert "Carbon dioxide" in text


def test_extracts_pptx_slides_and_tables():
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Respiratory system"
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1))
    table_shape.table.cell(0, 0).text = "Gas"
    table_shape.table.cell(0, 1).text = "Oxygen"
    stream = io.BytesIO()
    presentation.save(stream)

    units = document_service.extract_units(stream.getvalue(), "lesson.pptx")
    assert units[0]["source_label"] == "Slide 1"
    assert "Respiratory system" in units[0]["text"]
    assert "Oxygen" in units[0]["text"]


def test_extracts_xlsx_worksheet_cells():
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Photosynthesis"
    sheet.append(["Stage", "Location"])
    sheet.append(["Light reactions", "Thylakoid membrane"])
    stream = io.BytesIO()
    workbook.save(stream)

    units = document_service.extract_units(stream.getvalue(), "notes.xlsx")
    assert units[0]["source_label"] == "Sheet: Photosynthesis"
    assert "Light reactions" in units[0]["text"]
    assert "Thylakoid membrane" in units[0]["text"]